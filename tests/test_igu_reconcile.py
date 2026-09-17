"""Tests for scripts/igu_fleet.py + scripts/igu_reconcile.py — the IGU World LNG
Report ↔ backend intercomparison (docs/sops/igu_reconciliation.md).

The load-bearing guarantees: the extractor assigns words to columns from the header
positions and opens a row at ANY word in the IMO column (the orderbook's "Unknown"
rows included); the join is by IMO only; a field diff says which side moved
(`igu_changed` vs `backend_differs`); a backend `active` row that left the IGU fleet
is 'dropped'; and an IMO-less IGU row never pairs to a backend row. Pure logic, no
network, no PDF.
"""
from collections import Counter

import igu_fleet as ig
import igu_reconcile as ir


def _w(text, x0, top, h=8):
    return {"text": text, "x0": x0, "top": top, "bottom": top + h}


def _igu(imo, name, table="fleet", owner="MISC", builder="Mitsui", cap=130000,
         year=1995, vtype="Conventional", prop="Steam", cargo="Membrane"):
    return {"imo": imo, "imo_raw": imo or "Unknown", "name": name, "shipowner": owner,
            "shipbuilder": builder, "capacity": cap, "cargo_type": cargo,
            "vessel_type": vtype, "propulsion": prop, "delivery_year": year,
            "age": None, "table": table, "pdf_page": 1}


def _be(row, imo, name, status="active", owner="MISC", builder="Mitsui", cap="130000",
        year="1995", vtype="conventional", prop="steam", cargo="membrane",
        contract="", hull="", source="IGU"):
    return {"row_id": str(row), "sheet_row": row, "imo": imo, "name": name, "hull": hull,
            "status": status, "shipowner": owner, "shipbuilder": builder, "capacity": cap,
            "vessel_type": vtype, "propulsion": prop, "cargo_type": cargo,
            "delivery_year": year, "contract_date": contract, "original_source": source}


def _edition(fleet=(), orderbook=(), edition="2026", asof="end-2025"):
    return {"edition": edition, "asof": asof, "fleet": list(fleet),
            "orderbook": list(orderbook), "counts": {"fleet": len(fleet),
                                                     "orderbook": len(orderbook)}}


class TestExtractorRows:
    COLS = [("imo", "IMO Number", 10), ("name", "Name", 60), ("shipowner", "Shipowner", 140)]

    def _tbl(self, words):
        return {"columns": self.COLS, "words": words, "left": 0, "right": 300,
                "top": 0, "bottom": 500}

    def test_wrapped_cell_joins_the_row_above(self):
        rows = ig.read_table(self._tbl([
            _w("9030814", 10, 20), _w("Puteri", 60, 20), _w("Delima", 84, 20), _w("MISC", 140, 20),
            _w("9030826", 10, 40), _w("Puteri", 60, 40), _w("Nilam", 84, 40), _w("MISC", 140, 40),
            _w("Berhad", 140, 49),                       # wrapped owner line
        ]))
        assert [r["imo"] for r in rows] == ["9030814", "9030826"]
        assert rows[0]["name"] == "Puteri Delima"
        assert rows[1]["shipowner"] == "MISC Berhad"

    def test_unknown_imo_opens_a_row(self):
        rows = ig.read_table(self._tbl([
            _w("9030814", 10, 20), _w("Alpha", 60, 20), _w("MISC", 140, 20),
            _w("Unknown", 10, 40), _w("Hull", 60, 40), _w("8340", 80, 40), _w("BW", 140, 40),
        ]))
        assert len(rows) == 2
        assert rows[1]["imo"] == "Unknown" and rows[1]["name"] == "Hull 8340"

    def test_last_row_stops_at_a_gap(self):
        rows = ig.read_table(self._tbl([
            _w("9030814", 10, 20), _w("Alpha", 60, 20), _w("MISC", 140, 20),
            _w("Source:", 10.5, 200), _w("IGU", 60, 200),   # footer far below
        ]))
        assert rows[0]["name"] == "Alpha"

    def test_join_hyphen_wrap(self):
        assert ig._join(["Hudong (ex-", "Jiangnan)"]) == "Hudong (ex-Jiangnan)"

    def test_imo_check_digit(self):
        assert ig.imo_checksum_ok("9030814") is True
        assert ig.imo_checksum_ok("9030815") is False
        assert ig.imo_checksum_ok("Unknown") is False


class TestHelpers:
    def test_name_key_drops_nested_ex_tail_and_punctuation(self):
        assert ir._name_key("Energy Fortitude (ex-Victor Hugo (8107))") == ir._name_key("Energy Fortitude")
        assert ir._name_key("Taitar No. 1") == ir._name_key("Taitar No.1")

    def test_placeholders(self):
        for n in ("Hull 2541 (Hanwha)", "Hull H1884A", "Knutsen OAS - Hanwha - Dec 2025 - 1",
                  "Samsung (Seapeak 3)", "TBN", ""):
            assert ir.is_placeholder(n), n
        for n in ("Puteri Delima Satu", "BW Nivalis", "Trader III"):
            assert not ir.is_placeholder(n), n

    def test_contract_month_from_name(self):
        assert ir.contract_month_from_name("Knutsen OAS - Hanwha - Dec 2025 - 1") == "2025-12"
        assert ir.contract_month_from_name("Hull 8340") == ""

    def test_month_near(self):
        assert ir._month_near("2025-06", "2025-07")
        assert ir._month_near("2025-12", "2026-01")
        assert not ir._month_near("2025-06", "2025-08")
        assert not ir._month_near("", "2025-08")


class TestFieldDiffers:
    PAIRS = Counter({("mitsui", "mitsui e&s shipbuilding chiba"): 5})

    def test_blank_is_never_a_disagreement(self):
        assert not ir.field_differs("propulsion", "", "DFDE", self.PAIRS)
        assert not ir.field_differs("propulsion", "DFDE", None, self.PAIRS)

    def test_capacity_tolerance(self):
        assert not ir.field_differs("capacity", "174000", 174500, self.PAIRS)
        assert ir.field_differs("capacity", "200000", 174000, self.PAIRS)

    def test_igu_placeholder_name_is_not_a_diff(self):
        assert not ir.field_differs("name", "Al Sakhamah", "Hull 2559", self.PAIRS)
        assert ir.field_differs("name", "Hull 2559 (Hanwha)", "Al Sakhamah", self.PAIRS)

    def test_owner_token_overlap(self):
        assert not ir.field_differs("shipowner", "Maran Gas Maritime", "Maran Gas", self.PAIRS)
        assert ir.field_differs("shipowner", "BW", "Venture Global", self.PAIRS)

    def test_builder_learned_pairing(self):
        key = ir._key("Mitsui E&S Shipbuilding Chiba")
        pairs = Counter({("mitsui", key): ir.BUILDER_PAIR_MIN})
        assert not ir.field_differs("shipbuilder", "Mitsui E&S Shipbuilding Chiba", "Mitsui", pairs)
        assert ir.field_differs("shipbuilder", "Samsung Heavy Industries", "HD Hyundai", pairs)

    def test_case_insensitive_vocab(self):
        assert not ir.field_differs("vessel_type", "conventional", "Conventional", self.PAIRS)
        assert ir.field_differs("vessel_type", "conventional", "FSU", self.PAIRS)


class TestDiffKind:
    def _diff(self, be_name, cur_name, prev_name, have_prev=True):
        prev = _igu("9030814", prev_name) if prev_name else None
        d = ir.diff_fields(_igu("9030814", cur_name), prev, _be(25, "9030814", be_name),
                           Counter(), {}, have_prev=have_prev)
        return [x for x in d if x["field"] == "name"][0]["kind"]

    def test_igu_changed(self):
        assert self._diff("Golar Tundra", "Italis LNG (ex-Golar Tundra)", "Golar Tundra") == "igu_changed"

    def test_backend_differs(self):
        assert self._diff("Hoegh", "Hoegh Esperanza", "Hoegh Esperanza") == "backend_differs"

    def test_new_to_igu(self):
        assert self._diff("Karmol Antarctica", "Karadeniz LNGT Antarctica", None) == "new_to_igu"

    def test_no_prev_edition(self):
        assert self._diff("Hoegh", "Hoegh Esperanza", None, have_prev=False) == "no_prev_edition"

    def test_pending_agreement(self):
        pending = {"25": {"Name": [{"batch": "b1", "value": "Italis LNG", "decision": "accept",
                                    "confidence": "G"}]}}
        d = ir.diff_fields(_igu("9030814", "Italis LNG (ex-Golar Tundra)"),
                           _igu("9030814", "Golar Tundra"), _be(25, "9030814", "Golar Tundra"),
                           Counter(), pending)
        assert d[0]["pending_agrees"] is True


class TestStatusFinding:
    def test_delivered_per_igu_with_pending(self):
        pending = {"843": {"Status": [{"batch": "b1", "value": "active", "decision": "accept",
                                       "confidence": "G"}]}}
        sf = ir.status_finding(_igu("9970569", "Venture Pelican", year=2025),
                               _be(843, "9970569", "Hull 2541", status="on order", year="2026"),
                               pending, 2025)
        assert sf["finding"] == "delivered_per_igu" and sf["pending_agrees"]

    def test_on_order_at_cutoff_flags_year_conflict(self):
        sf = ir.status_finding(_igu("9968932", "BW Nivalis", table="orderbook", year=2026),
                               _be(787, "9968932", "BW Nivalis", year="2025"), {}, 2025)
        assert sf["finding"] == "on_order_at_igu_cutoff" and sf["delivery_year_conflict"]
        sf = ir.status_finding(_igu("1018676", "Celsius Georgetown", table="orderbook", year=2026),
                               _be(907, "1018676", "Celsius Georgetown", year="2026"), {}, 2025)
        assert sf["delivery_year_conflict"] is False

    def test_agreement_is_no_finding(self):
        assert ir.status_finding(_igu("9030814", "A"), _be(1, "9030814", "A"), {}, 2025) is None


class TestReconcileBuckets:
    def _run(self):
        cur = _edition(
            fleet=[_igu("9030826", "Puteri Nilam"),
                   _igu("9627497", "Maran Gas Efessos", owner="Maran Gas", builder="Hanwha Ocean",
                        cap=159800, year=2014),
                   _igu("9696266", "Hai Yang Shi You 301", vtype="Bunkering vessel", cap=30000)],
            orderbook=[_igu("", "Knutsen OAS - Hanwha - Dec 2025 - 1", table="orderbook",
                            owner="Knutsen OAS", builder="Hanwha Ocean", cap=174000, year=2029),
                       _igu("9961518", "H1884A", table="orderbook", owner="MOL"),
                       _igu("9961518", "Greenergy Wind", table="orderbook", owner="CNOOC")])
        prev = _edition(edition="2025", asof="end-2024",
                        fleet=[_igu("9030826", "Puteri Nilam"), _igu("9030814", "Puteri Delima")])
        backend = [
            _be(25, "9030814", "Puteri Delima"),                      # dropped
            _be(26, "9030826", "Puteri Nilam"),                       # matched, clean
            _be(61, "9211872", "Puteri Delima Satu", source="Clarkson"),   # active, never listed
            _be(912, "9961518", "Greenenergy Wind", status="on order", owner="MOL"),
            _be(1105, "", "Hanwha Ocean (Knutsen)", status="on order", owner="Knutsen OAS",
                builder="Hanwha Ocean", cap="174000", year="2029", contract="15-Dec-2025"),
        ]
        return ir.reconcile(cur, prev, backend, {})

    def test_dropped(self):
        rec = self._run()
        assert [x["backend"]["sheet_row"] for x in rec["dropped"]] == [25]

    def test_matched_by_imo_only(self):
        rec = self._run()
        assert sorted(m["backend"]["sheet_row"] for m in rec["matched"]) == [26, 912]

    def test_igu_only_candidate_and_out_of_scope(self):
        by = {x["igu"]["imo"]: x for x in self._run()["igu_only"]}
        assert by["9627497"]["out_of_scope"] is False
        assert by["9696266"]["out_of_scope"] is True

    def test_backend_not_in_igu_reasons(self):
        by = {x["backend"]["sheet_row"]: x["reason"] for x in self._run()["backend_not_in_igu"]}
        assert by[61].startswith("active and never listed")
        assert by[1105].startswith("no IMO")
        assert 25 not in by                                   # dropped rows are not double-counted

    def test_no_imo_rows_get_cluster_hints_never_pairs(self):
        rec = self._run()
        (c,) = rec["igu_no_imo"]
        assert c["igu_count"] == 1 and c["contract_month"] == "2025-12"
        assert [h["sheet_row"] for h in c["backend_hint_rows"]] == [1105]
        assert c["hint_strength"] == "strong"
        assert all(m["igu"]["imo"] for m in rec["matched"])

    def test_duplicates_kept_first_print_joined(self):
        rec = self._run()
        (d,) = rec["igu_duplicates"]
        assert d["imo"] == "9961518" and len(d["rows"]) == 2 and d["backend_rows"] == [912]
        m = [m for m in rec["matched"] if m["backend"]["sheet_row"] == 912][0]
        assert m["igu"]["name"] == "H1884A"

    def test_edition_diff(self):
        ed = self._run()["edition_diff"]
        assert [r["imo"] for r in ed["fleet_dropped"]] == ["9030814"]
        assert ed["fleet_dropped"][0]["backend_rows"] == [25]


class TestLeads:
    def test_load_leads_compacts_the_unit_record(self, tmp_path):
        import json
        p = tmp_path / "leads.json"
        p.write_text(json.dumps([
            {"imo": "9030814", "hits": [{"url": "https://www.shipvault.com/ships/1"}],
             "unit": {"unitid": 1, "unit": "TT", "name": "LIMA", "status": "SCRAPPED",
                      "owner": "BREAKERS", "fatedate": "2026-03-08T00:00:00", "delivered": None}},
            {"imo": "1040447", "hits": [], "unit": None}]))
        leads = ir.load_leads(p)
        assert leads == {"9030814": {"name": "TT LIMA", "status": "SCRAPPED", "owner": "BREAKERS",
                                     "fate_date": "2026-03-08", "delivered": "",
                                     "url": "https://www.shipvault.com/ships/1"}}

    def test_review_imos_skips_findings_a_pending_batch_covers(self):
        rec = {"dropped": [{"backend": {"imo": "9030814"}}],
               "matched": [{"backend": {"imo": "1"}, "status_finding": {"pending_agrees": True}},
                           {"backend": {"imo": "2"}, "status_finding": {"pending_agrees": False}},
                           {"backend": {"imo": "3"}, "status_finding": None}],
               "igu_only": [{"igu": {"imo": "4"}, "out_of_scope": False},
                            {"igu": {"imo": "5"}, "out_of_scope": True}],
               "backend_not_in_igu": [{"backend": {"imo": "6"}, "reason": "active and never listed"},
                                      {"backend": {"imo": ""}, "reason": "no IMO"}],
               "edition_diff": {"orderbook_dropped": [{"imo": "7"}]}}
        assert ir.review_imos(rec) == ["9030814", "2", "4", "6", "7"]


class TestWorkbook:
    def test_igu_mode_builds_every_sheet(self, tmp_path):
        import json
        from types import SimpleNamespace
        from openpyxl import load_workbook
        import build_workbook as bw
        rec = TestReconcileBuckets()._run()
        rec["leads"] = {"9030814": {"name": "TT LIMA", "status": "SCRAPPED", "owner": "BREAKERS",
                                    "fate_date": "2026-03-08", "delivered": "", "url": "u"}}
        rec["dropped"][0]["fate_vs_inclusion"] = "Dec 2025 or later — stays in scope"
        src = tmp_path / "igu_reconcile.json"
        src.write_text(json.dumps(rec))
        out = bw.build_igu(SimpleNamespace(reconcile=str(src), out=str(tmp_path) + "/"))
        wb = load_workbook(out)
        assert wb.sheetnames == ["README", "Summary", "Dropped_from_IGU", "Status_findings",
                                 "Field_diffs", "IGU_only", "Backend_not_in_IGU", "IGU_no_IMO",
                                 "IGU_duplicates", "Edition_diff", "QA_review"]
        dropped = [[c.value for c in r] for r in wb["Dropped_from_IGU"].iter_rows()]
        assert any(25 in r and "2026-03-08" in r for r in dropped)
