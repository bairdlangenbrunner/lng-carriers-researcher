"""
Build the review app's dataset (work/review_data.json) from one or more batch dirs.

Full values come from apply_batch's own item model (`_detect` + `_items_and_conflicts`),
never from decisions.csv, which truncates proposed_value / note to 80 characters. The
current decision per proposal comes from decisions.csv; backend context and current
values from the fresh pull (work/backend.csv). Every row a human sees is the LIVE sheet
row (`sheet_row_map`); `row_id` stays the key.

    python review_app/review_data.py --batches batches/<dir> [<dir> ...] [--out work/review_data.json]

Read-only over the batch dirs and the backend. Deterministic apart from `built`.
"""
import argparse
import csv
import json
import re
import sys
import time
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT / "scripts") not in sys.path:
    sys.path.insert(0, str(ROOT / "scripts"))

from apply_batch import _default_decision, _detect, _items_and_conflicts  # noqa: E402
from backend_io import load_backend  # noqa: E402
from paths import backend_csv_path, work_dir  # noqa: E402

ET = ZoneInfo("America/New_York")
DECISIONS = ("accept", "hold", "reject")

# Linked columns by exact backend header (plan fact 8); X <-> "X [ref]" is handled generically.
PAIRS = [("Name", "Other names"), ("Price", "Price currency"), ("Capacity", "Capacity units")]
# The Name <-> Other names prompt is not skippable in the UI (RF §4.16).
STRICT_PAIR = {"Name", "Other names"}
IGU_PDF = re.compile(r"igu-world-lng-report", re.I)


def read_csv(path):
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def split_urls(s):
    return [u.strip() for u in re.split(r",\s+|\s*\n\s*|\s+;\s+", s or "") if u.strip()]


def partner_columns(col, header):
    """Columns linked to `col` (fact 8), restricted to columns present in the header."""
    out = []
    for a, b in PAIRS:
        if col == a:
            out.append(b)
        elif col == b:
            out.append(a)
    if col.endswith(" [ref]"):
        out.append(col[:-6])
    elif f"{col} [ref]" in header:
        out.append(f"{col} [ref]")
    return [c for c in out if c in header]


def batch_info(batch_dirs, info_path=None):
    """dir -> {label, apply_order, applied} from an optional review_batches.json.

    Looked up at --info, else in any batches/*/review_batches.json that names one of the
    dirs. Absent: label = dir name, order = dir sort order, applied = a verify_report.csv
    exists in the dir (verify_apply.py only runs after an apply)."""
    info = {}
    paths = [Path(info_path)] if info_path else sorted((ROOT / "batches").glob("*/review_batches.json"))
    names = {d.name for d in batch_dirs}
    for p in paths:
        doc = json.loads(p.read_text(encoding="utf-8"))
        if names & set(doc):
            info = doc
            break
    out = {}
    for i, d in enumerate(sorted(batch_dirs, key=lambda d: d.name), 1):
        meta = info.get(d.name, {})
        out[d.name] = {
            "label": meta.get("label", d.name),
            "apply_order": meta.get("apply_order", i),
            "applied": bool(meta.get("applied", (d / "verify_report.csv").exists())),
        }
    return out


# ---- gate verdicts (plan fact 6) --------------------------------------------------

def _qa_sections(wb_path):
    """QA_review -> {section title: [row dicts]} (the build_combined.py approach)."""
    import openpyxl
    wb = openpyxl.load_workbook(wb_path, read_only=True)
    if "QA_review" not in wb.sheetnames:
        return {}
    out, title, header = defaultdict(list), "", None
    for r in wb["QA_review"].iter_rows(values_only=True):
        r = list(r)
        vals = [v for v in r if v not in (None, "")]
        if not vals:
            continue
        if len(vals) == 1 and r[0] not in (None, "") and not str(r[0]).isdigit():
            title, header = str(r[0]), None
            continue
        if header is None:
            header = [h for h in r if h not in (None, "")]
            continue
        out[title].append(dict(zip(header, r)))
    wb.close()
    return out


def _citation_ok(note):
    """A ref_fill Per-cell citation log note that records a pass: url_verifier's `: OK`, or an
    igu_refs pass written as "IGU <edition> report PDF … prints this hull for the row's IMO"
    (a note saying IGU prints *no* such column, or that an earlier pass was false, is a fail)."""
    if ": OK" in note:
        return True
    return bool(re.search(r"\bprints\b", note)) and not re.search(r"\b(not|no|false|nothing)\b", note)


def gate_verdicts(batch_dir, key_to_header):
    """(by_cell, by_url): (row_id, field, url) -> verdict and url -> verdict.

    The batch workbook's QA_review sheet, with gate_log.json (when the batch has one) laid
    over it — a gate log can cover only part of a batch (the roll-forward's covers 263 of
    473 checks). Nothing found -> None, and the UI shows "—" (a verdict is never invented)."""
    by_cell, by_url = {}, {}
    for wb in sorted(batch_dir.glob("lng_carrier_*.xlsx")):
        for title, rows in _qa_sections(wb).items():
            for q in rows:
                url = str(q.get("url") or "")
                if q.get("verdict") and url:                       # fix workbooks
                    by_cell[(str(q.get("row_id")), str(q.get("field")), url)] = str(q["verdict"])
                elif title == "Per-cell citation log" and url:     # ref_fill
                    note = str(q.get("note") or "")
                    field = key_to_header.get(str(q.get("field")), str(q.get("field")))
                    for u in split_urls(url):
                        by_cell[(str(q.get("row_id")), field, u)] = \
                            f"{'PASS' if _citation_ok(note) else 'FAIL'} ({note})"
                elif title == "URL verification log" and url:      # data_fill / discovery
                    res = q.get("result") if q.get("result") not in (None, "None") else q.get("status")
                    if res not in (None, "None", ""):
                        by_url[url] = f"{res} (URL log)"
    # a shipvault companion is ADDED only where the unit record corroborates the cell value
    # (shipvault_api_refs.py); an entry logged without a value was not gated
    sv = batch_dir / "shipvault_api_refs.json"
    if sv.exists():
        added = json.loads(sv.read_text(encoding="utf-8")).get("added")
        for a in added if isinstance(added, list) else []:
            rid, _, field = str(a.get("where", "")).partition("|")
            m = re.search(r"row_id (\S+?)\)", rid)
            rid, field = (m.group(1) if m else rid.strip()), field.strip()
            if a.get("result") == "ADDED" and a.get("value") and a.get("url"):
                field = field[:-6] if field.endswith(" [ref]") else field
                by_cell.setdefault((rid, field, a["url"]), "PASS (companion record corroborates)")
    gl = batch_dir / "gate_log.json"
    if gl.exists():
        for g in json.loads(gl.read_text(encoding="utf-8")):
            v = f"{'PASS' if g.get('ok') else 'FAIL'} ({g.get('reason', '')})"
            by_cell[(str(g.get("row_id")), g.get("field", ""), g.get("url", ""))] = v
    return by_cell, by_url


# ---- presentation: what the card shows first, what goes under "details" -------------

PROVENANCE_TAIL = re.compile(r"\s*\[([^\[\]]*(?:PDF p\.|IMO \d{7}|sole source)[^\[\]]*)\]\s*$")
PDF_PAGE = re.compile(r"PDF p\.\s*(\d+)")
SV_PAGE = re.compile(r"shipvault\.com/ships/(\d+)")
SV_API = re.compile(r"shipvaultapi[^/]*/api/units/(\d+)")
WHY_MAX = 240


def _balanced(head):
    """True when `head` does not end inside a parenthesis or a quotation."""
    opening = len(re.findall(r"(?:^|[\s(])'(?=\S)", head))
    closing = len(re.findall(r"(?<=\S)'(?=[\s.,;:)]|$)", head))
    return (head.count("(") == head.count(")") and head.count('"') % 2 == 0
            and head.count("“") == head.count("”") and opening == closing)


def split_note(note, current="", proposed="", labels=None):
    """(why, detail): the one sentence a reviewer needs, and the rest of the note.

    Nothing is dropped — `note` stays whole in the dataset and everything cut from `why`
    lands in `detail`, except a leading restatement of current -> proposed (the card shows
    those right above). Batch dir names become their short labels."""
    why, detail = (note or "").strip(), []
    for d, label in (labels or {}).items():
        why = why.replace(d, f"batch {label}")
    m = PROVENANCE_TAIL.search(why)
    if m:
        why, tail = why[:m.start()].rstrip(), m.group(1)
        detail.append(tail)
    lead = re.match(r"'([^']*)' -> '([^']*)':\s*", why)
    if lead and lead.group(1) == (current or "") and lead.group(2) == (proposed or ""):
        why = why[lead.end():]
    if len(why) > WHY_MAX:                  # a long closing parenthetical is an aside
        m = re.search(r"\s*\(([^()]{60,})\)\.?$", why)
        if m and m.start() >= 60:
            why, aside = why[:m.start()], m.group(1)
            detail.insert(0, aside)
    if len(why) > WHY_MAX:                  # else cut at the last clause boundary that fits
        cuts = [c for c in re.finditer(r"[.;]['\"]? (?=[A-Za-z'\"(])| \| | -- | — ", why[:WHY_MAX])
                if c.start() >= 60 and _balanced(why[:c.start() + 1])]
        if not cuts:                        # no boundary fits: end at the first whole sentence
            cuts = [c for c in re.finditer(r"\.['\"]? (?=[A-Z'\"(])", why[WHY_MAX:])
                    if _balanced(why[:WHY_MAX + c.start() + 1])][:1]
            cuts = [re.compile(re.escape(c.group())).search(why, WHY_MAX + c.start()) for c in cuts]
        if cuts:
            c = cuts[-1]
            head = why[:c.start() + 1] if why[c.start()] in ".;" else why[:c.start()]
            why, rest = head.rstrip(" ;"), why[c.end():].strip()
            detail.insert(0, rest)
    return why, detail


def ref_status(verdict):
    """verified / failed / read / unchecked (+ the gate's own reason for a failure)."""
    v = str(verdict or "")
    if not v:
        return "unchecked", ""
    if re.match(r"(PASS|OK|200)\b", v, re.I):
        return "verified", ""
    inner = re.sub(r"\s*\(URL log\)$", "", v)
    inner = re.sub(r"^\w+\s*\(?", "", inner).rstrip(")") if "(" in inner else ""
    if re.match(r"(FAIL|dead|banned|blocked)", v, re.I):
        return "failed", inner
    if re.match(r"READ", v, re.I):
        return "read", inner
    return "other", v


def ref_label(url, page=None):
    from urllib.parse import unquote, urlsplit
    if IGU_PDF.search(url):
        yr = re.search(r"report-(\d{4})", url, re.I)
        return "IGU World LNG Report" + (f" {yr.group(1)}" if yr else "") + (f", p.{page}" if page else "")
    if SV_PAGE.search(url) or SV_API.search(url):
        return "shipvault record"
    u = urlsplit(url)
    host = re.sub(r"^www\.", "", u.netloc)
    seg = [x for x in unquote(u.path).split("/") if x]
    slug = re.sub(r"\.(html?|php|aspx?)$", "", seg[-1]) if seg else ""
    slug = slug if len(slug) <= 64 else slug[:63] + "…"
    return host + (f" › {slug}" if slug and not slug.isdigit() else "") + (f", p.{page}" if page and url.lower().endswith(".pdf") else "")


def present_refs(refs, note):
    """Display refs: a label, a status, a #page link for a PDF the note locates, and a
    shipvault page merged with its companion unit record (one source, RF §6a.8)."""
    pm = PDF_PAGE.search(note or "")
    page = pm.group(1) if pm else None
    rank = {"verified": 0, "failed": 1, "read": 2, "other": 3, "unchecked": 4}
    out, by_unit = [], {}
    for r in refs:
        url = r["url"]
        status, reason = ref_status(r["verdict"])
        is_pdf = url.lower().split("#")[0].split("?")[0].endswith(".pdf")
        d = {"url": url, "href": url + (f"#page={page}" if page and is_pdf and "#" not in url else ""),
             "label": ref_label(url, page if is_pdf else None), "status": status, "reason": reason,
             "verdicts": [f"{url}: {r['verdict'] or 'not checked'}"]}
        m = SV_PAGE.search(url) or SV_API.search(url)
        if m and m.group(1) in by_unit:
            first = by_unit[m.group(1)]
            first["companion"] = url
            first["verdicts"] += d["verdicts"]
            if rank[status] < rank[first["status"]]:
                first["status"], first["reason"] = status, reason
            continue
        if m:
            by_unit[m.group(1)] = d
        out.append(d)
    return out


# ---- backend sync: what the live backend already holds ------------------------------

def _norm(s):
    return " ".join((s or "").split())


def _same(a, b):
    """Equal after whitespace normalisation, or as numbers (the sheet renders 165000000
    as `165000000.00`)."""
    a, b = _norm(a), _norm(b)
    if a == b:
        return True
    try:
        return float(a.replace(",", "")) == float(b.replace(",", ""))
    except ValueError:
        return False


def backend_state(kind, current, proposed, ref_urls, current_refs, append=False, keep_ref=False, ref_only=False):
    """"in_backend" when the pulled backend already holds the proposal (value and refs),
    "value_in_backend" when the value is there but a proposed ref is not, else "".

    Same test as verify_apply.py (whitespace-normalised equality); an appended element
    (`Other names`) has landed when every proposed element is in the cell."""
    have = {_norm(u) for u in current_refs}
    refs_in = all(_norm(u) in have for u in ref_urls)
    if kind == "ref" or ref_only:      # a ref-only line has landed when its refs are in the cell
        return "in_backend" if ref_urls and refs_in else ""
    if not _norm(proposed):
        return ""
    if append:
        cur = {_norm(x) for x in current.split(";")}
        value_in = all(_norm(x) in cur for x in proposed.split(";") if _norm(x))
    else:
        value_in = _same(current, proposed)
    if not value_in:
        return ""
    return "in_backend" if keep_ref or refs_in else "value_in_backend"


# ---- review items (plan fact 7) --------------------------------------------------

def collect_items(bdir, mode, srm, cell=None):
    """`cell(row_id, column)` reads the pulled backend; with it a conflict whose proposed value
    the backend now holds, and a duplicate pair with a row gone, carry `backend_resolved`."""
    name = bdir.name
    items = []

    def live(ids):
        return [srm[i] for i in ids if i in srm]

    cpath = bdir / "conflicts.csv"
    if cpath.exists():
        for i, c in enumerate(read_csv(cpath)):
            ids = [x for x in re.split(r"[/,;\s]+", c.get("row_id", "")) if x]
            rows = live(ids)
            if not rows:  # discovery flags carry "live row(s) 1168" in the text
                rows = [int(x) for x in re.findall(r"live row\(s\)\s*([\d,\s/-]+)", c.get("recommendation", ""))
                        for x in re.findall(r"\d+", x)][:20]
            items.append({
                "item_id": f"{name}::conflict:{i}", "type": "flag" if mode == "discovery" else "conflict",
                "batch": name, "live_rows": rows, "row_ids": ids,
                "title": " · ".join(x for x in (c.get("column"), c.get("backend_value") and
                                               f"{c['backend_value']} → {c.get('proposed_value', '')}") if x)
                or "backend flag",
                "detail": c.get("recommendation", ""), "urls": split_urls(c.get("sources", "")),
                # conflicts.csv rows carry no id: the write-back addresses record i and checks
                # row_id + column still match (apply_batch.py regenerates the file)
                "conflict_index": i, "conflict_decision": c.get("decision", ""),
                "conflict_match": [c.get("row_id", ""), c.get("column", "")],
                "backend_resolved": "the backend now holds the proposed value"
                if cell and len(ids) == 1 and _norm(c.get("proposed_value")) and
                _same(cell(ids[0], c.get("column", "")), c["proposed_value"]) else "",
            })
    mpath = bdir / "manual_review.json"
    if mpath.exists():
        for i, m in enumerate(json.loads(mpath.read_text(encoding="utf-8"))):
            rid = str(m.get("row_id", ""))
            title = m.get("name") or ""
            if m.get("field"):
                title = f"{m['field']}: {m.get('backend', '')} vs IGU {m.get('igu', '')}"
            elif m.get("verdict"):
                title = f"press follow-up: {m['verdict']}" + (f" ({m['name']})" if m.get("name") else "")
            urls = [u for k in ("status_urls", "year_urls", "name_urls") for u in (m.get(k) or [])]
            items.append({
                "item_id": f"{name}::manual:{i}", "type": "manual", "batch": name,
                "live_rows": live([rid]), "row_ids": [rid], "title": title or "manual review",
                "detail": " ".join(x for x in (m.get("why"), m.get("note")) if x), "urls": urls,
            })
    ppath = bdir / "proposed_review.json"
    if ppath.exists():
        for p in json.loads(ppath.read_text(encoding="utf-8")).get("programmes", []):
            urls = [e.get("url", "") for e in p.get("evidence", []) if e.get("url")]
            for a in p.get("row_actions", []):
                rid = str(a.get("row_id", ""))
                items.append({
                    "item_id": f"{name}::proposed:{rid}", "type": "proposed_bucket", "batch": name,
                    "live_rows": live([rid]), "row_ids": [rid],
                    "title": f"{p.get('programme', '')} · {a.get('action', '')}"
                             + ("" if rid in srm else " · row no longer in backend"),
                    "detail": a.get("note", ""), "urls": urls,
                })
    dpath = bdir / "dedupe_report.csv"
    if dpath.exists():
        for i, d in enumerate(read_csv(dpath)):
            ids = [x for x in re.split(r"[/,;\s]+", d.get("row_ids", "")) if x]
            items.append({
                "item_id": f"{name}::duplicate:{i}", "type": "duplicate", "batch": name,
                "live_rows": live(ids), "row_ids": ids,
                "title": f"{d.get('tier', '')} {d.get('reason', '')}".strip() or "possible duplicate",
                "detail": "; ".join(f"{k}: {v}" for k, v in d.items() if v and k not in ("row_ids",)),
                "urls": [],
                "backend_resolved": "a row of the pair is no longer in the backend"
                if len(ids) > 1 and len(live(ids)) < 2 else "",
            })
    return items


# ---- build -----------------------------------------------------------------------

def build(batch_dirs, backend_path=None, info_path=None):
    be = load_backend(backend_path)
    header, colmap = be.header, be.colmap
    H = {h: i for i, h in enumerate(header)}
    rows = be.row_by_id()
    srm = be.sheet_row_map()
    key_to_header = {k: header[v] for k, v in colmap.items()
                     if not k.startswith("_") and isinstance(v, int) and v < len(header)}
    info = batch_info(batch_dirs, info_path)

    def cell(rid, col):
        r = rows.get(rid)
        return r[H[col]].strip() if r is not None and col in H and len(r) > H[col] else ""

    # a discovery row is in the backend when a row with its IMO, Name or Hull number holds every
    # column the candidate fills (verify_apply.py's match) — over every row, since a stub pasted
    # without a row key is still the vessel, and a stub is not yet the row
    present = defaultdict(list)
    for r in be.data:
        for h in ("IMO number", "Name", "Hull number"):
            if h in H and len(r) > H[h] and _norm(r[H[h]]):
                present[(h, _norm(r[H[h]]))].append(r)

    def new_row_in_backend(rd):
        cols = [h for h, v in rd.items() if _norm(v) and h in H]
        return any(all(len(r) > H[h] and _norm(r[H[h]]) for h in cols)
                   for h in ("IMO number", "Name", "Hull number") if _norm(rd.get(h))
                   for r in present.get((h, _norm(rd.get(h))), []))

    batches, proposals, all_items = [], {}, []
    labels = {d: m["label"] for d, m in info.items() if m["label"] != d}
    for bdir in sorted(batch_dirs, key=lambda d: (info[d.name]["apply_order"], d.name)):
        mode, payload = _detect(bdir)
        items, _ = _items_and_conflicts(mode, payload, header, colmap)
        dec_path = bdir / "decisions.csv"
        dec = {r["id"]: (r.get("decision") or "").strip().lower() for r in read_csv(dec_path)} \
            if dec_path.exists() else {}
        by_cell, by_url = gate_verdicts(bdir, key_to_header)
        # fix-mode cell flags the item model does not carry (read from the already-loaded payload)
        fix_flags = {}
        if mode == "fix":
            for corr in payload.get("corrections", []):
                for c in corr.get("cells", []):
                    fix_flags[(str(corr["row_id"]), c.get("field", ""))] = c
        counts = {k: 0 for k in DECISIONS}
        meta = info[bdir.name]
        for it in items:
            key = f"{bdir.name}::{it['id']}"
            default = _default_decision(it["confidence"], it["derivable"])
            decision = dec.get(it["id"]) or default
            counts[decision] = counts.get(decision, 0) + 1
            rid, col = it["row_id"], it["column"]
            flags = []
            fc = fix_flags.get((rid, col), {})
            if fc.get("preserve_ref"):
                flags.append("preserve_ref")
            if fc.get("append_ref"):
                flags.append("append_ref")
            if it.get("ref_only"):
                flags.append("ref_only")
            if meta["applied"]:
                flags.append("applied")
            ref_urls = split_urls(it["ref_value"])
            if it["kind"] == "new_row":
                rd = it["row_data"] or {}
                ref_urls = list(dict.fromkeys(u for h, v in rd.items() if h.endswith("[ref]")
                                              for u in split_urls(v)))
            vfield = col[:-6] if it["kind"] == "ref" and col.endswith(" [ref]") else col
            refs = [{"url": u, "verdict": by_cell.get((rid, vfield, u)) or by_cell.get((rid, col, u))
                     or by_url.get(u)} for u in ref_urls]
            if fc.get("preserve_ref"):
                v = next((v for (r, f, _u), v in by_cell.items() if r == rid and f == col), None)
                if v:
                    flags.append("verdict:" + v)
            if ref_urls and all(IGU_PDF.search(u) for u in ref_urls):
                flags.append("igu_pdf_only")
            current_refs = split_urls(cell(rid, it["ref_column"] or (col if col.endswith("[ref]") else "")))
            if "preserve_ref" in flags and not current_refs:
                current_refs = split_urls(cell(rid, f"{col} [ref]"))
            proposed = it["value"] or it["ref_value"] if it["kind"] != "new_row" else ""
            if it["kind"] == "new_row":
                rd = it["row_data"] or {}
                state = "in_backend" if new_row_in_backend(rd) else ""
            else:
                state = backend_state(it["kind"], cell(rid, col), proposed, ref_urls, current_refs,
                                      append="append_ref" in flags, keep_ref="preserve_ref" in flags,
                                      ref_only="ref_only" in flags)
            if state:
                flags.append(state)
            why, detail = split_note(it["note"], cell(rid, col) if rid else "", proposed, labels)
            proposals[key] = {
                "why": why, "detail": detail, "sources": present_refs(refs, it["note"]),
                "batch": bdir.name, "id": it["id"], "kind": it["kind"],
                "row_id": rid, "cluster_id": it["cluster_id"], "column": col,
                "current": cell(rid, col) if rid else "",
                "proposed": proposed,
                "refs": refs, "ref_column": it["ref_column"], "current_refs": current_refs,
                "confidence": it["confidence"], "derivable": it["derivable"],
                # RF §5 rev 28: the grade is computed from what the §3.8c gate did, and
                # `why` is the sentence that says which route it took (or what held it).
                "confidence_why": it.get("confidence_why", ""),
                "prev_state": it["prev_state"], "default": default, "decision": decision,
                "note": it["note"], "links": [], "flags": flags,
                # a ref line: the backend value the ref is meant to cite
                "cited_value": cell(rid, vfield) if it["kind"] == "ref" and rid else "",
                "row_data": it["row_data"] if it["kind"] == "new_row" else None,
                "cluster_label": it.get("cluster_label", ""),
            }
        batches.append({"dir": bdir.name, "mode": mode, "label": meta["label"],
                        "apply_order": meta["apply_order"], "applied": meta["applied"],
                        "counts": counts, "n": len(items)})
        all_items.extend(collect_items(bdir, mode, srm, cell))

    # links: linked-column partners on the same row (any batch) + same id in another batch
    # keys compared canonically: an older batch keyed by legacy row_id and a newer one keyed by
    # UUID name the same vessel
    ck = be.canonical_key
    by_cell_key = defaultdict(list)
    for k, p in proposals.items():
        if p["row_id"]:
            by_cell_key[(ck(p["row_id"]), p["column"])].append(k)
    for k, p in proposals.items():
        if not p["row_id"]:
            continue
        links = []
        for other in by_cell_key[(ck(p["row_id"]), p["column"])]:
            if other != k:
                links.append(other)
                if "overlaps_batch" not in p["flags"]:
                    p["flags"].append("overlaps_batch")
        for pc in partner_columns(p["column"], header):
            links.extend(by_cell_key.get((ck(p["row_id"]), pc), []))
        if any(proposals[o]["column"] != p["column"] and {p["column"], proposals[o]["column"]} == STRICT_PAIR
               for o in links):
            p["flags"].append("strict_pair")
        p["links"] = links

    # vessels, ordered by live row (new discovery clusters last)
    order = {b["dir"]: (b["apply_order"], b["dir"]) for b in batches}
    grouped = defaultdict(list)
    for k, p in proposals.items():
        grouped[ck(p["row_id"]) if p["row_id"] else f"cluster:{p['batch']}:{p['cluster_id']}"].append(k)
    vessels = []
    for vid, keys in grouped.items():
        keys.sort(key=lambda k: (order[proposals[k]["batch"]], H.get(proposals[k]["column"], 999),
                                 proposals[k]["column"]))
        p0 = proposals[keys[0]]
        if p0["row_id"]:
            rid = ck(p0["row_id"])
            vessels.append({"row_id": rid, "live_row": srm.get(rid), "new": False,
                            "in_backend": rid in rows,
                            "name": cell(rid, "Name"), "imo": cell(rid, "IMO number"),
                            "status": cell(rid, "Status"), "shipbuilder": cell(rid, "Shipbuilder"),
                            "shipowner": cell(rid, "Shipowner"), "hull": cell(rid, "Hull number"),
                            "delivery_year": cell(rid, "Delivery year"), "proposals": keys})
        else:
            rd = p0["row_data"] or {}
            vessels.append({"row_id": "", "live_row": None, "new": True, "in_backend": False,
                            "cluster_id": p0["cluster_id"], "batch": p0["batch"],
                            "name": rd.get("Name") or p0["cluster_label"], "imo": rd.get("IMO number", ""),
                            "status": rd.get("Status", ""), "shipbuilder": rd.get("Shipbuilder", ""),
                            "shipowner": rd.get("Shipowner", ""), "hull": rd.get("Hull number", ""),
                            "delivery_year": rd.get("Delivery year", ""), "proposals": keys})
    vessels.sort(key=lambda v: (v["new"], v["live_row"] if v["live_row"] is not None else 10 ** 9,
                                v.get("cluster_id", ""), v["row_id"]))

    ordered = {k: proposals[k] for v in vessels for k in v["proposals"]}
    return {
        "built": datetime.now(ET).isoformat(timespec="seconds"),
        "backend_pulled": datetime.fromtimestamp(be.path.stat().st_mtime, ET).isoformat(timespec="seconds"),
        "header": header,
        "batches": batches,
        "vessels": vessels,
        "proposals": ordered,
        "items": all_items,
    }


def check_backend(backend_path):
    """Refuse without a pull; warn when the pull is older than a day. Returns the path."""
    p = Path(backend_path) if backend_path else backend_csv_path()
    if not p.exists():
        raise SystemExit(f"{p} not found — run `python scripts/pull_backend.py` first.")
    age_h = (time.time() - p.stat().st_mtime) / 3600
    if age_h > 24:
        print(f"WARNING: {p} is {age_h:.0f} h old — re-pull (python scripts/pull_backend.py) "
              "before deciding on current values.", file=sys.stderr)
    return p


def resolve_dirs(names):
    out = []
    for n in names:
        p = Path(n)
        if not p.is_dir():
            p = ROOT / "batches" / n
        if not p.is_dir():
            raise SystemExit(f"batch dir not found: {n}")
        out.append(p.resolve())
    return out


def write(data, out):
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    tmp = out.with_suffix(out.suffix + ".tmp")
    tmp.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    tmp.replace(out)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--batches", nargs="+", required=True)
    ap.add_argument("--out", default=None, help="default work/review_data.json")
    ap.add_argument("--backend", default=None)
    ap.add_argument("--info", default=None, help="review_batches.json (dir -> label / apply_order / applied)")
    args = ap.parse_args(argv)
    backend = check_backend(args.backend)
    data = build(resolve_dirs(args.batches), backend, args.info)
    out = args.out or (work_dir() / "review_data.json")
    write(data, out)
    tot = {k: sum(b["counts"].get(k, 0) for b in data["batches"]) for k in DECISIONS}
    print(f"review_data: {len(data['proposals'])} proposals ({tot['accept']} accept / {tot['hold']} hold / "
          f"{tot['reject']} reject), {len(data['vessels'])} vessels, {len(data['items'])} items -> {out}",
          file=sys.stderr)
    for b in data["batches"]:
        c = b["counts"]
        print(f"  {b['apply_order']:>3}  {b['label'][:40]:<40} {b['n']:>5}  accept {c['accept']:>4}  "
              f"hold {c['hold']:>4}  reject {c['reject']:>2}{'  (applied)' if b['applied'] else ''}",
              file=sys.stderr)
    return data


if __name__ == "__main__":
    main()
